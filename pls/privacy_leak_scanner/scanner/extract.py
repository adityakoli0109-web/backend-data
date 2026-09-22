from pathlib import Path
import csv
import json

from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document
from PIL import Image, ImageOps
import pytesseract

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
TEXT_EXTS = {".txt", ".md", ".log", ".json", ".xml", ".html", ".htm", ".yaml", ".yml"}


def _configure_tesseract() -> None:
    configured = getattr(pytesseract.pytesseract, "tesseract_cmd", "")
    if configured and configured != "tesseract":
        return
    common = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    for path in common:
        if path.exists():
            pytesseract.pytesseract.tesseract_cmd = str(path)
            return


def extract_image(path: Path) -> str:
    _configure_tesseract()
    with Image.open(path) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        return pytesseract.image_to_string(image)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if text.strip():
        return text

    # Scanned/image PDF fallback: render pages and OCR them.
    try:
        import fitz  # PyMuPDF
        _configure_tesseract()
        chunks = []
        with fitz.open(str(path)) as document:
            for page in document:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                chunks.append(pytesseract.image_to_string(image))
        return "\n".join(chunks)
    except Exception:
        return ""


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for ws in wb.worksheets:
        chunks.append(f"[sheet:{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            chunks.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(chunks)


def extract_csv(path: Path) -> str:
    try:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
            return "\n".join(" | ".join(row) for row in csv.reader(f))
    except Exception:
        return path.read_text(encoding="utf-8", errors="replace")


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
    return "\n".join(chunks)


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if ext == ".csv":
        return extract_csv(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".pptx":
        return extract_pptx(path)
    if ext in TEXT_EXTS:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if ext in IMAGE_EXTS:
        return extract_image(path)
    return ""
